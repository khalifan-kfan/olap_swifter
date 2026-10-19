"""
Full transfer and input files: Excel, Word and text inputs, long documents
extracted for the dataset concept, derived values and every value checked
against the source, EDA-shaped dummy data, the user's placement choice, the
private staging database kept out of what is shared, and data issues in the
results README. Runs on the AMR fixture schema.
"""

import csv
import json
import shutil
import zipfile
from pathlib import Path

import docx
import duckdb
import openpyxl
import pytest
import yaml
from typer.testing import CliRunner

from olap_swifter.paths import make_workspace, route, write_workspace_readme, zip_workspace
from olap_swifter.schema.checks import run_checks
from olap_swifter.schema.describe import describe
from olap_swifter.schema.extender import apply_patch
from olap_swifter.schema.materializer import WarehouseLoader
from olap_swifter.schema.profile import load_dictionary, load_eda
from olap_swifter.schema.readers import as_csv, data_entries, extract_document
from olap_swifter.schema.registry import SchemaRegistry
from olap_swifter.schema.synthetic import write_synthetic_dataset

SEED_DIR = Path(__file__).resolve().parent / "fixtures" / "schemas"
AMR = "amr_surveillance"


@pytest.fixture
def registry(tmp_path):
    d = tmp_path / "schemas"
    shutil.copytree(SEED_DIR, d)
    return SchemaRegistry(d)


def _rows(path):
    with open(path, newline="") as fh:
        return list(csv.reader(fh))


def _write_csv(path, header, rows):
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        w.writerows(rows)
    return path


def test_excel_sheets_are_data_files_and_load_like_csv(registry, tmp_path):
    dummy = write_synthetic_dataset(registry.get(AMR), str(tmp_path / "d"), n_visits=30)
    book = tmp_path / "data" / "amr.xlsx"
    book.parent.mkdir()
    wb = openpyxl.Workbook()
    wb.active.title = "results"
    for r in _rows(dummy["data"]):
        wb.active.append(r)
    wb.create_sheet("notes").append(["free text"])
    wb.save(book)

    assert data_entries(book.parent) == ["amr.xlsx#results", "amr.xlsx#notes"]  # workbook order
    with pytest.raises(ValueError, match="2 sheets"):
        as_csv(str(book))  # a workbook of several sheets needs one named
    staged = Path(as_csv(f"{book}#results"))
    assert staged.parent.name == ".sheets" and _rows(staged) == _rows(dummy["data"])

    report = WarehouseLoader(duckdb.connect()).load(registry.get(AMR), f"{book}#results", "amr_xlsx",
                                                    is_synthetic=True)
    assert report["loaded"] and report["postload"]["all_passed"], report["postload"]["failures"]


def test_word_and_excel_catalogues_and_excel_eda(tmp_path):
    d = docx.Document()
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Version", "2"  # not a dictionary: skipped
    t = d.add_table(rows=1, cols=3)
    for i, h in enumerate(["Variable", "Description", "Type"]):
        t.rows[0].cells[i].text = h
    for name, desc, typ in [("visit_id", "Visit", "string"), ("mic", "MIC", "number")]:
        cells = t.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = name, desc, typ
    d.save(tmp_path / "dict.docx")
    cat = load_dictionary(str(tmp_path / "dict.docx"))
    assert [v["name"] for v in cat["variables"]] == ["visit_id", "mic"]
    assert cat["variables"][1]["type"] == "DOUBLE"

    wb = openpyxl.Workbook()
    for r in [["variable", "type", "missing_pct", "distinct", "categories"],
              ["ast_interpretation", "VARCHAR", "0", "3", "S:40|R:25|I:5"]]:
        wb.active.append(r)
    wb.save(tmp_path / "amr.xlsx")
    eda = load_eda(str(tmp_path / "amr.xlsx"))
    assert eda["variables"][0]["stats"]["categories"] == {"S": 40, "R": 25, "I": 5}


def test_long_documents_are_extracted_and_linked_to_the_concept(registry, tmp_path, monkeypatch):
    monkeypatch.setenv("OLAP_SWIFTER_WORKSPACE", str(tmp_path / "ws"))
    make_workspace("amr_site")
    folder = route("inputs") / "amr_site"
    dummy = write_synthetic_dataset(registry.get(AMR), str(tmp_path / "d"), n_visits=20)
    shutil.copy(dummy["data"], folder / "data" / "isolates.csv")
    shutil.copy(dummy["catalogue"], folder / "catalogue" / "isolates.csv")
    filler = "routine laboratory surveillance text " * 130
    pages = [f"Background population setting. {filler}", f"The primary objective is resistance. {filler}",
             f"Confidential identifiers are removed. {filler}", filler, filler]
    (folder / "documents" / "protocol.txt").write_text("\n".join(pages))
    (folder / "documents" / "short.md").write_text("one page")
    (folder / "documents" / "scan.rtf").write_text("{\\rtf1}")

    out = describe("amr_site", study="s", main_question="q", one_row_is="r", main_event="e", quiet=True)
    docs = {d["file"]: d for d in out["documents"]}
    long = docs["documents/protocol.txt"]
    assert long["pages"] > 3 and Path(long["text"]).exists()
    assert "===== page 2 =====" in Path(long["text"]).read_text()
    assert 2 in long["concept_links"]["study.main_question"] and 3 in long["concept_links"]["privacy"]
    assert "text" not in docs["documents/short.md"]  # three pages or fewer: read as it is
    assert out["unsupported_files"] == {"documents": ["scan.rtf"]}
    with pytest.raises(ValueError, match="placement"):
        describe("amr_site", quiet=True, placement="merge")


def test_an_underivable_value_stops_the_load(registry, tmp_path):
    """A derived value that cannot be computed from inputs that are present would be lost: a transfer error."""
    spec = apply_patch(registry.get(AMR), {"reason": "birth year", "add_columns": {"dim_patient": {
        "birth_year": {"type": "BIGINT", "role": "attribute", "expr": 'year(TRY_CAST("date_of_birth" AS DATE))',
                       "inputs": ["date_of_birth"]}}}}, registry)
    data = _write_csv(tmp_path / "p.csv", ["visit_id", "patient_id", "facility_id", "visit_date", "date_of_birth"],
                      [["V1", "P1", "F1", "2025-01-02", "1980-04-01"], ["V2", "P2", "F1", "2025-01-03", "UNKNOWN"],
                       ["V3", "P3", "F1", "2025-01-04", ""]])
    report = WarehouseLoader(duckdb.connect()).load(spec, str(data), "dob")
    failed = {(f["check"], f["target"]): f["actual"] for f in report["preflight"]["failures"]}
    assert failed == {("derived_completeness", "dim_patient.birth_year"): 1}  # the blank one is simply missing
    assert report["loaded"] is False


def test_value_parity_catches_a_changed_value(registry, tmp_path):
    conn = duckdb.connect()
    dummy = write_synthetic_dataset(registry.get(AMR), str(tmp_path / "d"), n_visits=30)
    WarehouseLoader(conn).load(registry.get(AMR), dummy["data"], "amr_dummy", is_synthetic=True)
    clean = run_checks(conn, registry.get(AMR), "amr_dummy")["postload"]
    assert clean["all_passed"], clean["failures"]
    conn.execute("UPDATE fact_ast_results SET mic_ug_ml = mic_ug_ml + 1 "
                 "WHERE _source_row = (SELECT MIN(_source_row) FROM fact_ast_results WHERE mic_ug_ml IS NOT NULL)")
    failed = {(f["check"], f["target"]) for f in run_checks(conn, registry.get(AMR), "amr_dummy")["postload"]["failures"]}
    assert ("value_parity", "fact_ast_results") in failed


def test_dummy_data_follows_the_dataset_eda(registry, tmp_path):
    spec = registry.get(AMR)
    eda = {"row_count": 1000, "variables": [
        {"name": "mic_ug_ml", "type": "DOUBLE", "stats": {"null_pct": 80.0, "min": 2, "max": 8}},
        {"name": "ast_interpretation", "type": "VARCHAR", "stats": {"null_pct": 0, "categories": {"R": 90, "S": 10}}}]}
    (tmp_path / "eda.yaml").write_text(yaml.safe_dump(eda))
    spec.datasets["amr_site"] = {"eda": str(tmp_path / "eda.yaml")}
    out = write_synthetic_dataset(spec, str(tmp_path / "d"), n_visits=150)
    assert out["shaped_by_eda_variables"] == 2
    rows = list(csv.DictReader(open(out["data"])))
    mic = [r["mic_ug_ml"] for r in rows]
    present = [float(v) for v in mic if v]
    assert 0.6 < 1 - len(present) / len(mic) and all(2 <= v <= 8 for v in present)  # EDA null rate and range
    sir = [r["ast_interpretation"] for r in rows if r["ast_interpretation"]]
    assert sir.count("R") > 3 * sir.count("S")  # EDA proportions, inside the allowed codes


def test_extending_or_reusing_tables_follows_the_users_placement(registry, tmp_path, monkeypatch):
    from olap_swifter.schema_cli import schema_app
    monkeypatch.setenv("OLAP_SWIFTER_SCHEMAS_DIR", str(registry.dir))
    concept = tmp_path / "dataset_concept.yaml"
    concept.write_text(yaml.safe_dump({"placement": "own_mart"}))
    mart = registry.get(AMR).to_yaml_dict()
    mart["schema_id"] = "lab_mart"
    path = tmp_path / "lab_mart.yaml"
    path.write_text(yaml.safe_dump(mart))
    cat = _write_csv(tmp_path / "cat.csv", ["variable"], [["visit_id"]])
    args = ["--dataset", "lab", "--catalogue", str(cat), "--eda", str(cat), "--concept", str(concept), "--json"]

    res = CliRunner().invoke(schema_app, ["register", str(path)] + args)
    assert res.exit_code != 0 and "own data mart" in res.output and "dim_visit" in res.output
    res = CliRunner().invoke(schema_app, ["attach", AMR] + args)
    assert res.exit_code != 0 and "user's decision" in res.output


def test_staging_is_never_shared(tmp_path, monkeypatch):
    from olap_swifter.feedback.db_viewer import find_databases
    monkeypatch.setenv("OLAP_SWIFTER_WORKSPACE", str(tmp_path / "ws"))
    make_workspace()
    for name in ("warehouse.duckdb", "warehouse_staging.duckdb"):
        duckdb.connect(str(route("warehouses") / name)).close()
    names = zipfile.ZipFile(zip_workspace()["zip"]).namelist()
    assert "ws/warehouse/warehouse.duckdb" in names and not any("staging" in n for n in names)
    assert [p.name for p in find_databases()] == ["warehouse.duckdb"]


def test_data_issues_are_listed_in_the_readme(registry, tmp_path, monkeypatch):
    monkeypatch.setenv("OLAP_SWIFTER_WORKSPACE", str(tmp_path / "ws"))
    make_workspace()
    data = _write_csv(tmp_path / "s.csv", ["visit_id", "patient_id", "facility_id", "visit_date", "sex"],
                      [["V1", "P1", "F1", "2025-01-02", "Female"], ["V2", "P1", "F1", "2025-01-03", "Male"],
                       ["V3", "P2", "F1", "2025-01-04", "female"]])
    conn = duckdb.connect(str(route("warehouse")))
    report = WarehouseLoader(conn).load(registry.get(AMR), str(data), "amr_site")
    assert report["loaded"] and report["preflight"]["all_passed"]  # data issues never block
    rec = json.loads(Path(report["data_issues"]["file"]).read_text())
    assert {i["check"] for i in rec["issues"]} == {"entity_consistency", "allowed_values"}
    text = write_workspace_readme().read_text()
    section = text.split("## Data issues for cleaning")[1].split("## Token weight")[0]
    assert "| entity_consistency | `dim_patient` |" in section and "`dim_patient.gender`" in section


def test_document_extraction_without_pages_estimates_them(tmp_path):
    d = docx.Document()
    for _ in range(8):
        d.add_paragraph("study design population setting " * 60)
    d.save(tmp_path / "p.docx")
    info = extract_document(tmp_path / "p.docx", tmp_path / "out")
    assert info["pages"] >= 4 and Path(info["text"]).exists() and "study.description" in info["concept_links"]
